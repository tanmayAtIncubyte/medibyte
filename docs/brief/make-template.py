"""Build docs/QA-ASSESSMENT-TEMPLATE.xlsx, the workbook manual QA candidates fill in.

It mirrors "What to submit" in docs/CANDIDATE-BRIEF.md: traditional test cases,
BDD scenarios and bug reports, each marked positive or negative. The example
rows describe a generic email app on purpose, so they can't hint at anything in
MediByte. Run from the repo root (needs openpyxl):

    python3 -m venv /tmp/xlsx && /tmp/xlsx/bin/pip install openpyxl
    /tmp/xlsx/bin/python docs/brief/make-template.py
"""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

OUT = Path(__file__).resolve().parents[1] / "QA-ASSESSMENT-TEMPLATE.xlsx"
ROWS = 300  # pre-formatted rows per sheet; dropdowns reach further

# Brand tokens from docs/brief/CANDIDATE-BRIEF.print.html
FOREST, MINT, MUTED, BORDER = "014D43", "EBFFF6", "66716D", "D5E6DC"
EXAMPLE_FILL = "F3F5F4"

HEAD_FONT = Font(bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor=FOREST)
EX_FONT = Font(italic=True, color=MUTED)
EX_FILL = PatternFill("solid", fgColor=EXAMPLE_FILL)
LINE = Side(style="thin", color=BORDER)
GRID = Border(left=LINE, right=LINE, top=LINE, bottom=LINE)
WRAP_TOP = Alignment(wrap_text=True, vertical="top")
EX_NOTE = "Example from a different product. Delete this row before you submit."

TYPE = ["Positive", "Negative"]
STATUS = ["Pass", "Fail", "Blocked", "Not run"]
SEVERITY = ["Critical", "High", "Medium", "Low"]
CATEGORY = ["Functional", "Accessibility", "Performance", "Security", "UI", "UX"]

SHEETS = [
    {
        "title": "Test cases",
        "columns": [
            ("TC ID", 10), ("Area / page", 16), ("Title", 32), ("Type", 11),
            ("Preconditions", 24), ("Test data", 24), ("Steps", 40),
            ("Expected result", 32), ("Actual result", 32), ("Status", 11), ("Bug ID", 10),
        ],
        "lists": {"Type": TYPE, "Status": STATUS},
        "examples": [
            ["TC-EX-1", "Compose", "Send an email with a subject and body", "Positive",
             "Signed in to the email app",
             "To: a valid address\nSubject: Incubyte\nBody: QA test for Incubyte",
             "1. Click Compose\n2. Enter the recipient, subject and body\n3. Click Send",
             '"Message sent" appears and the email is in Sent with the same subject and body',
             "As expected", "Pass", ""],
            ["TC-EX-2", "Compose", "Block sending to an invalid address", "Negative",
             "Signed in to the email app",
             "To: not-an-email\nSubject: Incubyte",
             "1. Click Compose\n2. Enter the recipient and subject\n3. Click Send",
             "Sending is blocked and the invalid address is flagged",
             '"Message sent" appears and the email is in Sent', "Fail", "BUG-EX-1"],
        ],
    },
    {
        "title": "BDD scenarios",
        "columns": [
            ("Scenario ID", 12), ("Feature", 18), ("Scenario", 32), ("Type", 11),
            ("Given", 34), ("When", 38), ("Then", 38), ("TC ID", 10),
        ],
        "lists": {"Type": TYPE},
        "examples": [
            ["BDD-EX-1", "Compose email", "Send an email with a subject and body", "Positive",
             "Given I am signed in\nAnd I have opened a new message",
             'When I enter a valid recipient\nAnd I enter the subject "Incubyte"\n'
             'And I enter the body "QA test for Incubyte"\nAnd I click Send',
             'Then I see "Message sent"\nAnd the email is in Sent with the subject "Incubyte"',
             "TC-EX-1"],
            ["BDD-EX-2", "Compose email", "Block sending to an invalid address", "Negative",
             "Given I am signed in\nAnd I have opened a new message",
             'When I enter the recipient "not-an-email"\nAnd I click Send',
             "Then sending is blocked\nAnd the invalid address is flagged",
             "TC-EX-2"],
        ],
    },
    {
        "title": "Bug reports",
        "columns": [
            ("Bug ID", 10), ("Title", 34), ("Area / page", 16), ("Category", 14),
            ("Severity", 11), ("Environment", 20), ("Preconditions", 22),
            ("Steps to reproduce", 40), ("Expected result", 30), ("Actual result", 30),
            ("Evidence", 26), ("TC ID", 10),
        ],
        "lists": {"Category": CATEGORY, "Severity": SEVERITY},
        "examples": [
            ["BUG-EX-1", "Email is sent to an invalid address without a warning", "Compose",
             "Functional", "Medium", "Chrome, macOS", "Signed in",
             "1. Click Compose\n2. In To, enter not-an-email\n3. Enter the subject Incubyte\n"
             "4. Click Send",
             "Sending is blocked and the invalid address is flagged",
             '"Message sent" appears and the email is in Sent',
             "screenshots/BUG-EX-1.png", "TC-EX-2"],
        ],
    },
]

INSTRUCTIONS = [
    "Use one row per test case, scenario or bug. Add as many rows as you need.",
    'Write your test cases two ways: traditional on the "Test cases" sheet, and '
    'Given / When / Then on the "BDD scenarios" sheet.',
    "Cover both positive cases (it works as intended) and negative cases (invalid input, "
    "edge cases, error paths). Mark each one in the Type column.",
    'Log every defect you find on the "Bug reports" sheet. Link bugs and test cases to '
    "each other by ID.",
    "Save screenshots or recordings in your repository and put the file path or link in "
    "the Evidence column.",
    "The grey rows on each sheet are examples from a different product. Delete them "
    "before you submit.",
    "Commit this workbook to a Git repository and send us the repository URL, as the "
    "brief explains.",
]

SEVERITY_GUIDE = [
    ("Critical", "Blocks a core task, or puts data, money or safety at risk. No workaround."),
    ("High", "A major feature gives wrong results or fails. Any workaround is hard to find."),
    ("Medium", "Wrong or confusing behaviour that has a reasonable workaround."),
    ("Low", "Cosmetic or minor. It doesn't stop the task."),
]


def build_start(ws):
    ws.title = "Start here"
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 96

    ws["A1"] = "MediByte QA assessment: submission template"
    ws["A1"].font = Font(bold=True, size=16, color=FOREST)
    ws["A2"] = "Fill in this workbook, commit it to a Git repository and send us the URL."
    ws["A2"].font = Font(color=MUTED)

    row = 4
    for label in ["Your name", "Date submitted", "Browser(s) and OS", "Git repository URL"]:
        ws.cell(row, 1, label).font = Font(bold=True)
        cell = ws.cell(row, 2)
        cell.fill = PatternFill("solid", fgColor=MINT)
        cell.border = GRID
        row += 1

    row += 1
    ws.cell(row, 1, "How to fill this in").font = Font(bold=True, size=13, color=FOREST)
    row += 1
    for n, text in enumerate(INSTRUCTIONS, start=1):
        ws.cell(row, 1, n).alignment = Alignment(horizontal="left", vertical="top")
        ws.cell(row, 2, text).alignment = WRAP_TOP
        row += 1

    row += 1
    ws.cell(row, 1, "Severity guide").font = Font(bold=True, size=13, color=FOREST)
    row += 1
    for level, meaning in SEVERITY_GUIDE:
        ws.cell(row, 1, level).font = Font(bold=True)
        ws.cell(row, 2, meaning).alignment = WRAP_TOP
        row += 1


def build_sheet(wb, spec):
    ws = wb.create_sheet(spec["title"])
    ws.sheet_properties.tabColor = FOREST
    names = [name for name, _ in spec["columns"]]

    for col, (name, width) in enumerate(spec["columns"], start=1):
        letter = ws.cell(1, col).column_letter
        ws.column_dimensions[letter].width = width
        head = ws.cell(1, col, name)
        head.font, head.fill, head.border = HEAD_FONT, HEAD_FILL, GRID
        head.alignment = Alignment(wrap_text=True, vertical="center")
        for r in range(2, ROWS + 2):
            body = ws.cell(r, col)
            body.alignment, body.border = WRAP_TOP, GRID
    ws.row_dimensions[1].height = 30

    for r, example in enumerate(spec["examples"], start=2):
        for col, value in enumerate(example, start=1):
            cell = ws.cell(r, col, value or None)
            cell.font, cell.fill = EX_FONT, EX_FILL
        ws.cell(r, 1).comment = Comment(EX_NOTE, "Incubyte")
        # Excel won't grow a row for wrapped text it didn't type itself, so size it here.
        lines = max(
            sum(-(-len(part) // int(width * 1.1)) or 1 for part in str(value).split("\n"))
            for value, (_, width) in zip(example, spec["columns"])
        )
        ws.row_dimensions[r].height = 15 * lines + 4

    for column, options in spec["lists"].items():
        letter = ws.cell(1, names.index(column) + 1).column_letter
        dv = DataValidation(
            type="list", formula1='"' + ",".join(options) + '"', allow_blank=True,
            showErrorMessage=True, errorTitle=column,
            error="Pick one of: " + ", ".join(options),
        )
        dv.add(f"{letter}2:{letter}1000")
        ws.add_data_validation(dv)

    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{ws.cell(1, len(names)).column_letter}{ROWS + 1}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "1:1"


def main():
    wb = Workbook()
    build_start(wb.active)
    for spec in SHEETS:
        build_sheet(wb, spec)
    wb.properties.title = "MediByte QA assessment: submission template"
    wb.properties.creator = "Incubyte"
    wb.properties.lastModifiedBy = "Incubyte"
    wb.save(OUT)
    print(f"wrote {OUT.relative_to(Path.cwd()) if OUT.is_relative_to(Path.cwd()) else OUT}")


if __name__ == "__main__":
    main()
